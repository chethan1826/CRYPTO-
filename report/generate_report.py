from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "DES_Secure_Document_Encryption_Report.docx"
NAVY = "1C3557"
BLUE = "315B83"
PALE = "EAF0F5"
PALE2 = "F5F7FA"
INK = "222B35"
MUTED = "5B6875"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=85, start=105, bottom=85, end=105):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, *, bold=False, color=INK, size=9.1):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.02
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def set_keep(paragraph, keep=True):
    paragraph.paragraph_format.keep_together = keep


def add_body(doc, text, *, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Inches(0.22)
    p.paragraph_format.widow_control = True
    if bold_lead and text.startswith(bold_lead):
        lead, rest = text.split(" ", 1)
        p.add_run(lead + " ").bold = True
        p.add_run(rest)
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.27 + 0.18 * level)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.04
    p.add_run(text)
    return p


def add_number(doc, text):
    number = getattr(add_number, "counter", 0) + 1
    add_number.counter = number
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.27)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.04
    p.add_run(f"{number}. ").bold = True
    p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    if widths:
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, title in enumerate(headers):
        set_cell_text(header.cells[i], title, bold=True, color="FFFFFF", size=font_size)
        shade(header.cells[i], NAVY)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        if widths:
            for cell, width in zip(cells, widths):
                cell.width = Inches(width)
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=font_size)
            if row_idx % 2 == 1:
                shade(cells[i], PALE2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.32)
    sec.footer_distance = Inches(0.34)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True
    for name, size in (("Heading 1", 15), ("Heading 2", 11.2), ("Heading 3", 10.4)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(4 if name == "Heading 1" else 5)
        style.paragraph_format.space_after = Pt(5 if name == "Heading 1" else 3)
        style.paragraph_format.keep_with_next = True
    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(23)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(10.1)
        style.font.color.rgb = RGBColor.from_string(INK)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    r = footer.add_run("DES Document Encryption System  |  Page ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    add_field(footer, "PAGE")

    return doc


def build():
    doc = setup_document()

    # Page 1: academic cover page, following the supplied report's structure.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("VISVESVARAYA TECHNOLOGICAL UNIVERSITY, BELAGAVI")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    r = p.add_run("JSS ACADEMY OF TECHNICAL EDUCATION, BENGALURU")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("CRYPTOGRAPHY AND NETWORK SECURITY (BCS703)")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("ASSIGNMENT 1\nIMPLEMENTATION OF A CRYPTOGRAPHIC ALGORITHM")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(5)
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)
    p.add_run("DES Secure Document Encryption System")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("A browser-based demonstration of DES file encryption and decryption")
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run("Submitted by").bold = True
    details = doc.add_table(rows=2, cols=2)
    details.style = "Table Grid"
    details.autofit = False
    details.columns[0].width = Inches(1.35)
    details.columns[1].width = Inches(4.1)
    for i, (label, value) in enumerate((("Student Name", "[Enter your full name]"), ("USN", "[Enter your university seat number]"))):
        set_cell_text(details.rows[i].cells[0], label, bold=True, color=NAVY, size=10)
        set_cell_text(details.rows[i].cells[1], value, size=10)
        shade(details.rows[i].cells[0], PALE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.add_run("Guided by").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.add_run("Dr. Pooja H")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Associate Professor, Department of CSE")
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    p.add_run("Academic Year 2026-2027")

    page_break(doc)

    # Page 2: static contents map, aligned to the requested pagination.
    add_heading(doc, "TABLE OF CONTENTS", 1)
    entries = [
        ("Abstract", "3"),
        ("1. Introduction", "3"),
        ("2. Problem Statement", "4"),
        ("3. Objectives", "4"),
        ("4. Scope of the Project", "5"),
        ("5. Technologies and Requirements", "6"),
        ("6. System Architecture", "7"),
        ("7. Algorithms Used", "8"),
        ("8. Implementation Methodology", "9"),
        ("9. Testing and Results", "10"),
        ("10. Security Considerations and Limitations", "11"),
        ("11. Future Enhancements", "12"),
        ("12. Conclusion", "13"),
        ("13. References", "14"),
    ]
    for title, page in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.08)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.25), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.add_run(title)
        p.add_run("\t" + page)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.left_indent = Inches(0.08)
    r = p.add_run("Project note: This application is a local academic demonstration of a legacy cipher. It is not intended to protect confidential documents.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    page_break(doc)

    # Page 3: Abstract and introduction.
    add_heading(doc, "ABSTRACT", 1)
    add_body(doc, "This project presents a browser-based demonstration of the Data Encryption Standard (DES) for encrypting and decrypting documents. A user selects a file, supplies an eight-character printable ASCII key, and requests encryption from a local Java service. The service uses Java Cryptography Architecture and Java Cryptography Extension (JCA/JCE) to apply DES in CBC mode with PKCS#5 padding. A fresh random 8-byte initialization vector (IV) is generated for every encryption. The downloaded .des container stores a format marker, the IV, and the ciphertext. Decryption reads the same container and uses the same key to recover the original bytes.")
    add_body(doc, "The application binds its HTTP server to the local loopback address, keeps file bytes in request memory, and does not save uploaded plaintext to a database or server-side file. Account passwords are hashed with bcrypt on the Java server; only email addresses and bcrypt hashes are persisted in a local properties file. The implementation uses the JDK cryptographic provider for DES; its rounds are not manually reimplemented. The code compiles and the local service starts, while end-to-end file round-trip and login results must be recorded after a manual demonstration. DES is obsolete because its 56-bit effective key is too short for modern protection, and CBC encryption alone does not authenticate ciphertext. The system is intended only for coursework and algorithm study.")
    add_heading(doc, "1. INTRODUCTION", 1)
    add_body(doc, "Cryptography transforms information so that it can be read only in an authorized form. In symmetric encryption, the sender and receiver use the same secret key for encryption and decryption. The readable input is plaintext; the transformed output is ciphertext. DES is a historically important symmetric block cipher that is included in the course topics and is useful for explaining block processing, key use, modes of operation, and security limitations.")
    add_body(doc, "This project applies DES to a practical file workflow. A user creates a local account or signs in, chooses a document in the browser, enters a key, and downloads the encrypted result. To recover the file, the user selects the .des output and supplies the same key. The file wrapper carries the random IV needed by CBC mode. Java's standard cryptographic API performs the transformation, while the surrounding application validates the account and file requests, transfers bytes, and offers downloads.")
    add_body(doc, "The term secure in the project name describes the academic encryption workflow, not a claim that DES is safe for real data. The application makes this distinction visible and explains why modern deployments should choose authenticated encryption such as AES-GCM [1, 5].")

    page_break(doc)

    # Page 4: problem and objectives.
    add_heading(doc, "2. PROBLEM STATEMENT", 1)
    add_body(doc, "Students can learn the name and basic theory of a block cipher without seeing how an actual document moves through an encryption and decryption workflow. A useful demonstration should let a student choose a real file, apply a symmetric-key algorithm, save a ciphertext artifact, and recover the file with the matching key. It should also explain the parameters needed to reverse the operation, such as the initialization vector and padding.")
    add_body(doc, "The problem addressed by this project is to build a small, understandable local application that demonstrates those steps with DES, while clearly stating that DES is deprecated and that CBC does not provide ciphertext authentication. The application also needs a basic local account flow that protects the encryption workspace without storing plaintext passwords. It must reject malformed key inputs, identify its own encrypted-file format, and avoid persisting uploaded plaintext on the server.")
    add_heading(doc, "3. OBJECTIVES", 1)
    for text in [
        "Demonstrate symmetric encryption and decryption of arbitrary document bytes using DES.",
        "Use an explicit JCA/JCE transformation: DES/CBC/PKCS5Padding.",
        "Generate a new random 8-byte IV for each encryption and include it in the downloaded container.",
        "Provide a browser workflow for file selection, key entry, encryption, decryption, and downloads.",
        "Validate that the key is exactly eight printable ASCII characters and limit the input size.",
        "Protect the encryption workspace with local registration/login and server-side bcrypt password hashing.",
        "Describe DES key-size limitations, CBC behavior, and the absence of integrity protection.",
        "Keep the service local and process file bytes in memory without database or cloud storage.",
    ]:
        add_bullet(doc, text)

    page_break(doc)

    # Page 5: scope.
    add_heading(doc, "4. SCOPE OF THE PROJECT", 1)
    add_body(doc, "The project is a single-machine demonstration. A browser presents a file-encryption interface and sends binary file data to a Java HTTP server listening on 127.0.0.1:8080. The server applies DES through JCA/JCE, returns the encrypted container or decrypted bytes, and closes the request. The browser initiates the corresponding download.")
    add_heading(doc, "4.1 Features included", 2)
    for text in [
        "Encryption and decryption of non-empty files up to 25 MiB per plaintext input.",
        "An eight-character printable ASCII DES key, supplied by the user and reused for decryption.",
        "DES in CBC mode with PKCS#5 padding and a random 8-byte IV.",
        "A simple version marker and IV prefix in the .des file format.",
        "Local request handling, basic validation, generic decryption error responses, and browser downloads.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "4.2 Features outside the scope", 2)
    for text in [
        "Remote identity providers, multi-factor authentication, password recovery, remote sharing, cloud storage, or long-term key management.",
        "A hand-written DES round function, key-schedule visualizer, cryptanalysis tool, or formal security proof.",
        "Integrity or authenticity verification for ciphertext, sender identity, or tamper-proof transfer.",
        "A production service exposed beyond the local computer or a guarantee for confidential documents.",
    ]:
        add_bullet(doc, text)
    add_body(doc, "The interface requests Google Fonts for typography when internet access is available. That stylesheet request is separate from the file-processing API; passwords, document bytes, and the DES key are sent to the local Java service, not to a cloud storage provider. The authentication and encryption features have no cloud dependency.")

    page_break(doc)

    # Page 6: technologies and requirements.
    add_heading(doc, "5. TECHNOLOGIES AND REQUIREMENTS", 1)
    add_table(doc, ["Component", "Technology", "Purpose in this project"], [
        ("Backend language", "Java 11 or newer", "Runs the local HTTP server and file-processing logic."),
        ("HTTP server", "JDK HttpServer", "Serves the UI and handles POST requests on loopback."),
        ("Cryptography", "JCA/JCE Cipher", "Applies DES/CBC/PKCS5Padding using the installed provider."),
        ("Randomness", "java.security.SecureRandom", "Generates the CBC initialization vector."),
        ("Password hashing", "jBCrypt 0.4", "Stores bcrypt password hashes; plaintext passwords are not persisted."),
        ("Frontend", "HTML, CSS, JavaScript", "Selects files and keys and downloads results."),
        ("Storage", "Local properties file and browser downloads", "Stores account hashes; does not persist uploaded documents."),
    ], widths=[1.25, 1.9, 3.15], font_size=8.6)
    add_heading(doc, "5.1 Functional requirements", 2)
    add_body(doc, "The application must register and authenticate local accounts, store only bcrypt password hashes, issue an expiring session cookie, and require an authenticated session for encryption and decryption. It must accept file bytes and the DES key, encrypt the file, return a .des container, accept that container for decryption, and return recovered bytes. It must reject invalid key formats and empty or oversized requests.")
    add_heading(doc, "5.2 Software and hardware requirements", 2)
    add_body(doc, "Software: JDK 11 or later, the included jBCrypt 0.4 jar, Windows PowerShell to run start.ps1, and a current web browser. Maven or a database is not required. Hardware: an ordinary computer capable of running a JDK and browser, with sufficient memory for the file being processed. The server reads the request into memory and limits plaintext uploads to 25 MiB; this is a demonstration limit rather than a streaming design.")

    page_break(doc)

    # Page 7: architecture.
    add_heading(doc, "6. SYSTEM ARCHITECTURE", 1)
    add_body(doc, "The system has a browser client and a local Java service. The browser handles account forms, file selection, key entry, requests, and downloads. The Java service verifies passwords with jBCrypt, issues an HttpOnly session cookie, checks authorization, and invokes the JCA/JCE provider for DES. Both the interface and API are served by the same Java process.")
    diagram = doc.add_table(rows=1, cols=3)
    diagram.style = "Table Grid"
    diagram.autofit = False
    for cell, width in zip(diagram.rows[0].cells, [2.0, 2.0, 2.0]):
        cell.width = Inches(width)
    blocks = [
        ("BROWSER CLIENT", "HTML, CSS, JavaScript\nFile picker and key input\nDownload response bytes"),
        ("LOCAL JAVA SERVICE", "HttpServer on 127.0.0.1:8080\nPOST /api/register and /login\nPOST /api/encrypt and /decrypt"),
        ("CRYPTO AND PASSWORD HASHING", "jBCrypt cost 12\nDES/CBC/PKCS5Padding\nSecureRandom IV"),
    ]
    for cell, (title, detail) in zip(diagram.rows[0].cells, blocks):
        cell.text = ""
        shade(cell, PALE)
        cell_margins(cell, top=130, bottom=130, start=110, end=110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(NAVY)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(detail)
        run.font.name = "Times New Roman"
        run.font.size = Pt(8.4)
        run.font.color.rgb = RGBColor.from_string(INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("Request: file bytes + X-Des-Key  |  Response: encrypted or recovered bytes")
    r.italic = True
    r.font.size = Pt(8.8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    add_heading(doc, "6.1 Encryption flow", 2)
    add_number.counter = 0
    for text in [
        "The browser sends a POST request to /api/encrypt with the file bytes as the body and the key in the X-Des-Key header.",
        "The server checks the request method, payload length, and key format, then generates an 8-byte random IV.",
        "JCA/JCE encrypts the bytes. The server returns the marker, IV, and ciphertext as an octet-stream response.",
        "JavaScript turns the response into a downloadable file with the .des suffix.",
    ]:
        add_number(doc, text)
    add_heading(doc, "6.2 Decryption flow", 2)
    add_body(doc, "The browser posts the .des bytes and the same key to /api/decrypt. The service first checks the session, then checks the marker, reads the IV, decrypts the ciphertext, and returns the recovered bytes. Uploaded documents are not persisted; bcrypt account hashes are saved locally.")

    page_break(doc)

    # Page 8: algorithm.
    add_heading(doc, "7. ALGORITHMS USED", 1)
    add_heading(doc, "7.1 Data Encryption Standard", 2)
    add_body(doc, "DES is a symmetric block cipher that operates on 64-bit blocks. Its supplied key is 64 bits wide, but eight bits are parity bits, leaving 56 effective key bits. The standard DES transformation applies an initial permutation, then 16 Feistel rounds. In each round, the right half is expanded, combined with a round subkey, passed through substitution boxes, and permuted. The two halves are exchanged according to the Feistel structure; a final permutation completes the block transformation. Decryption uses the same structure with the round keys applied in reverse order.")
    add_body(doc, "This application does not implement those internal steps itself. It requests the named JCA/JCE transformation DES/CBC/PKCS5Padding from the Java provider [3]. The provider performs DES and applies the selected mode and padding.")
    add_heading(doc, "7.2 Cipher Block Chaining", 2)
    add_body(doc, "CBC links each plaintext block to the preceding ciphertext block. For the first block, the IV takes the place of a preceding ciphertext block. For 64-bit plaintext blocks P_i and ciphertext blocks C_i, encryption is C_1 = E_K(P_1 XOR IV), and C_i = E_K(P_i XOR C_(i-1)) for i greater than 1. Decryption reverses the block transformation and XORs the result with the IV or preceding ciphertext block. The IV is not a secret key, but it should be unpredictable and must be available during decryption [2].")
    add_heading(doc, "7.3 Padding and file container", 2)
    add_body(doc, "DES processes complete 8-byte blocks. PKCS#5 padding adds bytes so that the final plaintext block is complete, including a full padding block when the input length is already a multiple of eight. The Java provider removes this padding during successful decryption. The app's file container is: 8-byte ASCII marker DESFILE1, followed by the 8-byte IV, followed by the padded ciphertext. The marker identifies this app's format; it is not a message authentication code and does not prove that the file is genuine or unchanged.")
    add_heading(doc, "7.4 BCrypt password hashing", 2)
    add_body(doc, "The account feature uses jBCrypt to hash passwords on the server with a cost factor of 12. BCrypt incorporates a random salt in each stored hash and is deliberately slower than a fast general-purpose hash. During login, the library checks the submitted password against the stored bcrypt string; the original password is not stored. BCrypt is a one-way password hash, not reversible encryption. The demonstration limits passwords to 72 UTF-8 bytes, matching bcrypt's input limit.")

    page_break(doc)

    # Page 9: methodology.
    add_heading(doc, "8. IMPLEMENTATION METHODOLOGY", 1)
    add_body(doc, "The implementation is divided into a browser interface and a Java request handler. The account screen posts email and password fields to the local service; after registration or a successful login, the server issues an HttpOnly, SameSite=Strict session cookie that expires after one hour. The encryption and decryption endpoints require that session. For file processing, the browser posts binary data directly and downloads the response as application/octet-stream.")
    add_heading(doc, "8.1 Key and request validation", 2)
    add_body(doc, "The UI enables encryption only when a file is selected and the key field contains eight characters. The backend performs its own validation: the key must be exactly eight printable ASCII bytes. The server reads at most the configured input limit plus one byte to detect an oversized request. Plaintext encryption input is limited to 25 MiB; the decryption request limit includes a small allowance for the container marker, IV, and padding.")
    add_heading(doc, "8.2 Encryption operation", 2)
    add_number.counter = 0
    for text in [
        "Generate an 8-byte IV using SecureRandom for each encryption operation.",
        "Create a Cipher using the explicit transformation DES/CBC/PKCS5Padding.",
        "Initialize the cipher with the DES key and IV, then encrypt the in-memory file bytes.",
        "Construct the output as DESFILE1 || IV || ciphertext and return it to the browser.",
    ]:
        add_number(doc, text)
    add_heading(doc, "8.3 Decryption operation and runtime", 2)
    add_body(doc, "The server verifies the marker and minimum container length, extracts the IV and ciphertext, and initializes the same transformation for decryption. If the key or content is invalid, the server returns a generic failure response; because CBC has no authentication tag, a wrong key or modification is not guaranteed to be detected. The application starts by compiling DesServer.java and launching the Java class through start.ps1. The server binds only to the loopback interface at 127.0.0.1:8080.")

    page_break(doc)

    # Page 10: testing/results with strictly qualified evidence.
    add_heading(doc, "9. TESTING AND RESULTS", 1)
    add_body(doc, "A build check was completed for the Java backend with JDK 25, and the local server was launched successfully. Account registration/login and the file encryption/decryption round trip have not been independently exercised while preparing this report. The table separates confirmed build checks from demonstrations that should be run and recorded before submission.")
    add_table(doc, ["Check", "Expected result", "Evidence / status"], [
        ("Compile backend", "DesServer.java compiles into the build/classes directory.", "PASS - compilation succeeded with JDK 25."),
        ("Start local service", "Console reports http://localhost:8080 and the server serves the UI.", "PASS - service startup message was observed."),
        ("Encrypt and decrypt a PDF", "A .des file is produced; decryption with the same key returns byte-identical PDF data.", "NOT RUN - perform with a sample PDF and record the observed result."),
        ("Wrong key or modified ciphertext", "The service may reject bad padding, but CBC does not guarantee detection of all changes.", "NOT RUN - do not treat an error as an integrity guarantee."),
        ("Invalid key and oversized input", "Invalid key returns HTTP 400; an oversized body returns HTTP 413.", "Source behavior is implemented; live checks are not run."),
    ], widths=[1.2, 2.8, 2.25], font_size=8.2)
    add_heading(doc, "9.1 Manual PDF demonstration procedure", 2)
    add_number.counter = 0
    for text in [
        "Start the application and open http://localhost:8080 in a browser.",
        "Choose a small PDF and enter an eight-character printable ASCII key, for example Study123.",
        "Encrypt the PDF and download the resulting filename.pdf.des file.",
        "Select that .des file, enter the identical key, decrypt, and download the recovered PDF.",
        "Open the recovered PDF and compare its bytes or a file hash with the original; record the actual outcome and screenshots.",
    ]:
        add_number(doc, text)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    r = p.add_run("Result to complete after demonstration: [Record PDF name, key handling method, file sizes, comparison method, and observed pass/fail result. Do not include the secret key in screenshots or the submitted report.]")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    page_break(doc)

    # Page 11: limitations.
    add_heading(doc, "10. SECURITY CONSIDERATIONS AND LIMITATIONS", 1)
    add_heading(doc, "10.1 DES is obsolete", 2)
    add_body(doc, "DES has only 56 effective key bits. Exhaustive key search is practical with modern resources, so DES must not protect real documents. NIST withdrew FIPS 46-3 in 2005 because DES no longer provided the required security [1]. Its role here is historical and instructional.")
    add_heading(doc, "10.2 CBC provides no authentication", 2)
    add_body(doc, "The application uses CBC for confidentiality but stores no authentication tag or message authentication code. CBC ciphertext can be modified, and successful padding removal does not establish that the plaintext is correct. The marker only identifies a file layout. NIST distinguishes confidentiality-only modes from authenticated modes and recommends considering authentication where feasible [6].")
    add_heading(doc, "10.3 Key, IV, and transport handling", 2)
    add_body(doc, "The key is entered by the user and sent in the X-Des-Key HTTP header to a server bound to loopback. Account passwords are hashed with bcrypt and only hashes are persisted in build/users.properties. Sessions are held in memory and represented by an HttpOnly, SameSite=Strict cookie. This remains a local classroom demonstration: it has no TLS, rate limiting, password recovery, or remote access-control model. The DES key is present in browser and server memory during the request. The IV is generated randomly and stored alongside the ciphertext; it is not secret.")
    add_heading(doc, "10.4 Data handling and operational limits", 2)
    add_body(doc, "The service does not write file contents to a database or persistent server-side file, but request and response data occupy memory while processed. The current implementation reads the whole file into memory and therefore is not suited to large files or concurrent production traffic. The output does not contain the original filename as authenticated metadata; the browser chooses the download name. A production design would need authenticated encryption, careful key management, safe metadata, transport protection, and operational controls.")

    page_break(doc)

    # Page 12: future work.
    add_heading(doc, "11. FUTURE ENHANCEMENTS", 1)
    add_body(doc, "The project can be improved by keeping DES as a labelled legacy demonstration while adding a separate modern mode for practical file protection. Changes should be introduced as distinct workflows so students can compare the algorithms without confusing DES with current security guidance.")
    for title, detail in [
        ("Add AES-GCM as a modern mode.", "Use a supported AES key size, a unique nonce, and the authentication tag to protect both confidentiality and integrity. NIST specifies GCM as authenticated encryption [4, 5]."),
        ("Improve key management.", "Generate a random key for each file and protect it with an appropriate key-wrapping design, or derive a key from a password using a dedicated password-based key derivation function and a stored salt. Do not reuse short human-chosen text as a production key."),
        ("Add integrity and format metadata.", "Version the container and authenticate its header, nonce, ciphertext, original file metadata, and any recipient information."),
        ("Support streaming files.", "Process larger documents as streams with strict size limits and resource cleanup rather than holding complete input and output copies in memory."),
        ("Expand verification.", "Add known-answer vectors for algorithm components, round-trip checks for several file types, wrong-key and tampering cases, size-boundary checks, and reproducible performance measurements."),
        ("Improve the academic demonstration.", "Add a block diagram, downloadable test evidence, optional step-by-step DES structure explanation, and a comparison of DES and AES security, performance, block size, and authentication support."),
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.27)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(title + " ")
        r.bold = True
        p.add_run(detail)

    page_break(doc)

    # Page 13: conclusion.
    add_heading(doc, "12. CONCLUSION", 1)
    add_body(doc, "The DES Secure Document Encryption System demonstrates a local account flow and a file-encryption workflow: a user registers or signs in, selects a document, provides a symmetric key, downloads a ciphertext container, and uses the same key to request recovery of the original bytes. Registration uses server-side bcrypt password hashing, while the Java backend uses DES/CBC/PKCS5Padding with a random 8-byte IV. Uploaded file data is processed in memory rather than saved by the application.")
    add_body(doc, "The implementation is useful for connecting course concepts to a practical interface, but it does not make DES safe for real-world use. DES was withdrawn by NIST, and CBC encryption in this application does not authenticate the ciphertext. The Java provider performs the cipher operation; the project does not implement the internal DES rounds by hand. Compilation and local service startup were confirmed, while the PDF round-trip and other functional cases remain to be demonstrated and documented.")
    add_body(doc, "The project meets its educational goal when its limitations are explained alongside the demonstration. A future practical version should use an authenticated modern cipher such as AES-GCM, supported by proper key handling and tests. The current app should remain a coursework example only.")

    page_break(doc)

    # Page 14: references.
    add_heading(doc, "13. REFERENCES", 1)
    references = [
        "[1] National Institute of Standards and Technology (NIST), “Withdrawal of FIPS 46-3, Data Encryption Standard (DES),” May 19, 2005. https://csrc.nist.gov/news/2005/withdrawal-of-fips-46-3-fips-74-and-fips-81 (accessed 26 October 2026).",
        "[2] NIST, FIPS PUB 81, DES Modes of Operation, 1980. https://csrc.nist.gov/files/pubs/fips/81/final/docs/fips81.pdf (accessed 26 October 2026).",
        "[3] Oracle, Java Cryptography Architecture Reference Guide, JDK 22. https://docs.oracle.com/en/java/javase/22/security/java-cryptography-architecture-jca-reference-guide.html (accessed 26 October 2026).",
        "[4] NIST, FIPS PUB 197, Advanced Encryption Standard (AES), updated 9 May 2023. https://csrc.nist.gov/pubs/fips/197/final (accessed 26 October 2026).",
        "[5] NIST, SP 800-38D, Recommendation for Block Cipher Modes of Operation: Galois/Counter Mode (GCM) and GMAC, 2007. https://csrc.nist.gov/pubs/sp/800/38/d/final (accessed 26 October 2026).",
        "[6] NIST, Block Cipher Techniques: Authentication for Confidentiality Modes. https://csrc.nist.gov/projects/block-cipher-techniques/bcm/authentication-for-confidentiality-modes (accessed 26 October 2026).",
        "[7] Project implementation, DesServer.java, README.md, and project.md, Secure Document Encryption System, accessed 26 October 2026.",
        "[8] Damien Miller, jBCrypt, Java implementation of OpenBSD's Blowfish password hashing code, version 0.4. https://github.com/jeremyh/jBCrypt (accessed 27 October 2026).",
    ]
    for ref in references:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(9)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(ref)
        run.font.size = Pt(9.2)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run("Report prepared for the DES document-encryption project. Replace the student details on the cover page and complete the pending test results after performing the demonstration.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    doc.core_properties.title = "DES Secure Document Encryption System Report"
    doc.core_properties.subject = "Cryptography and Network Security assignment report"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "DES, file encryption, cryptography, CBC, JCA, JCE"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
